import json
import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest

BASE_DIR = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BASE_DIR / "backend" / "src"))

from openpyxl import load_workbook  # noqa: E402

from editor import LOG_SHEET, apply_edits, convert_value  # noqa: E402
from inspector import inspect_workbook  # noqa: E402
from rules import validate_rules  # noqa: E402
from screening import RulesError, run_screening  # noqa: E402

TEST_FILE = BASE_DIR / "data" / "test_all_checks.xlsx"
RULES_FILE = BASE_DIR / "data" / "study_rules.json"


@pytest.fixture
def rules():
    return json.loads(RULES_FILE.read_text(encoding="utf-8"))


def test_planted_issues_match_expected(rules):
    result = run_screening(TEST_FILE.read_bytes(), rules)
    expected = pd.read_excel(TEST_FILE, sheet_name="Expected_Issues", dtype=str)

    actual = set(zip(result["issues"]["Cell"], result["issues"]["Issue_Type"]))
    assert actual == set(zip(expected["Cell"], expected["Expected_Issue"]))
    assert result["warnings"] == []


def test_report_contains_all_sheets(rules):
    result = run_screening(TEST_FILE.read_bytes(), rules)
    book = pd.ExcelFile(pd.io.common.BytesIO(result["report"]))
    assert book.sheet_names == ["Summary", "Issues", "Rules", "Raw_Data"]


def test_missing_rule_column_warns(rules):
    rules["date_columns"] = {"Visit_date": {"min": "2026-09-01"}}
    result = run_screening(TEST_FILE.read_bytes(), rules)
    assert any("Visit_date" in w for w in result["warnings"])


@pytest.mark.parametrize("broken, message", [
    ({"numeric_columns": {"Age": {"min": 18}}}, "min과 max"),
    ({"numeric_columns": {"Age": {"min": 100, "max": 18}}}, "보다 큽니다"),
    ({"cross_checks": [{"left": "SBP", "op": "=>", "right": "DBP"}]}, "연산자"),
    ({"id_column": {"column": "Subject_ID", "pattern": "SUBJ-("}}, "정규식"),
    ({"date_columns": {"Visit_Date": {"min": "2026-13-01"}}}, "날짜 형식"),
    ({"categorical_column": {"Sex": ["M", "F"]}}, "알 수 없는 항목"),
])
def test_invalid_rules_are_rejected(broken, message):
    assert any(message in e for e in validate_rules(broken))
    with pytest.raises(RulesError):
        run_screening(TEST_FILE.read_bytes(), broken)


def test_inspect_suggests_column_types():
    info = inspect_workbook(TEST_FILE.read_bytes())
    types = {c["column"]: c["type"] for c in info["columns"]}

    assert info["sheets"] == ["Raw_Data", "Expected_Issues"]
    assert types["SBP"] == "numeric"
    assert types["Visit_Date"] == "date"
    assert types["Sex"] == "categorical"
    assert types["Notes"] == "empty"


def test_edits_fix_issues_and_are_logged(rules):
    edits = [
        {"cell": "G5", "value": "170", "issue": "SUSPECTED_TYPO"},
        {"cell": "H4", "value": "72.5", "issue": "EXTRA_WHITESPACE"},
        {"cell": "C20", "value": "2026-10-01", "issue": "DATE_OUT_OF_PERIOD"},
    ]
    edited = apply_edits(TEST_FILE.read_bytes(), "Raw_Data", edits, "2026-10-23 10:00:00")

    fixed_cells = {e["cell"] for e in edits}
    before = run_screening(TEST_FILE.read_bytes(), rules)["issues"]
    after = run_screening(edited, rules)["issues"]
    assert fixed_cells <= set(before["Cell"])
    assert not fixed_cells & set(after["Cell"])
    assert len(after) == len(before) - len(edits)

    book = load_workbook(BytesIO(edited))
    assert book.sheetnames == ["Raw_Data", "Expected_Issues", LOG_SHEET]
    assert book["Raw_Data"]["G5"].value == 170
    assert book["Raw_Data"]["C20"].value == "2026-10-01"

    log = [list(row) for row in book[LOG_SHEET].iter_rows(min_row=2, values_only=True)]
    assert log[0] == ["2026-10-23 10:00:00", "Raw_Data", "G5", "Height_cm", "17O", "170", "SUSPECTED_TYPO"]
    assert [row[2] for row in log] == ["G5", "H4", "C20"]


def test_edits_append_to_existing_log():
    first = apply_edits(TEST_FILE.read_bytes(), "Raw_Data", [{"cell": "G5", "value": "170"}])
    second = apply_edits(first, "Raw_Data", [{"cell": "I7", "value": "150"}])

    log = load_workbook(BytesIO(second))[LOG_SHEET]
    assert [row[2] for row in log.iter_rows(min_row=2, values_only=True)] == ["G5", "I7"]


def test_header_row_cannot_be_edited():
    with pytest.raises(ValueError, match="머리글"):
        apply_edits(TEST_FILE.read_bytes(), "Raw_Data", [{"cell": "N1", "value": "x"}])


@pytest.mark.parametrize("text, before, expected", [
    ("120", "12O", 120),
    ("72.5", " 72.5", 72.5),
    ("-3", None, -3),
    ("001", "01", "001"),
    ("SUBJ-016", "subj-016", "SUBJ-016"),
    ("", "x", None),
    ("2026-10-01", datetime(2026, 12, 1), datetime(2026, 10, 1)),
    ("2026-10-01", "2027-03-01", "2026-10-01"),
])
def test_convert_value(text, before, expected):
    assert convert_value(text, before) == expected


def test_table_matches_issue_cells(rules):
    result = run_screening(TEST_FILE.read_bytes(), rules)
    table = result["table"]

    assert len(table["rows"]) == result["summary"]["total_rows"]
    assert table["columns"][6] == "Height_cm"
    # 문제 셀 G5 = 7번째 컬럼, 엑셀 5행 = rows[3]
    assert table["rows"][3][6] == "17O"
    for issue in result["issues"].itertuples():
        if issue.Excel_Row > 1:
            col = table["columns"].index(issue.Column)
            assert table["rows"][issue.Excel_Row - 2][col] == issue.Original_Value
