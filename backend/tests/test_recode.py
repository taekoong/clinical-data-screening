from io import BytesIO
from pathlib import Path
import sys

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from editor import apply_edits, preview_recode


def workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    for row in [("ID", "Sex"), ("001", "f"), ("002", "F"), ("003", "Male"),
                ("004", "X"), ("005", " f "), ("006", None), ("007", "f")]:
        sheet.append(row)
    book.create_sheet("Other")["A1"] = "unchanged"
    sheet["B2"].number_format = "@"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_preview_and_apply_exact_matches_preserve_other_values_and_history():
    original = workbook()
    preview = preview_recode(original, "Data", "Sex", "f", "F")
    assert [edit["cell"] for edit in preview["edits"]] == ["B2", "B8"]
    source = load_workbook(BytesIO(original))
    assert source["Data"]["B2"].value == "f"
    assert "Change_Log" not in source.sheetnames

    edited = apply_edits(original, "Data", preview["edits"], "2026-10-23 15:00:00")
    book = load_workbook(BytesIO(edited))
    assert [book["Data"].cell(row, 2).value for row in range(2, 9)] == ["F", "F", "Male", "X", " f ", None, "F"]
    assert book["Other"]["A1"].value == "unchanged"
    assert book["Data"]["B2"].number_format == "@"
    log = list(book["Change_Log"].values)
    assert len(log) == 3
    assert log[1] == ("2026-10-23 15:00:00", "Data", "B2", "Sex", "f", "F", "CATEGORY_RECODE")
    second = preview_recode(edited, "Data", "Sex", "Male", "M")
    again = load_workbook(BytesIO(apply_edits(edited, "Data", second["edits"])))
    assert again["Data"]["B4"].value == "M"
    assert again["Change_Log"].max_row == 4


@pytest.mark.parametrize("target", ["001", "1", "=1+1"])
def test_category_targets_are_literal_text(target):
    data = workbook()
    preview = preview_recode(data, "Data", "Sex", "f", target)
    book = load_workbook(BytesIO(apply_edits(data, "Data", preview["edits"])))
    assert book["Data"]["B2"].value == target
    assert book["Data"]["B2"].data_type == "s"
    assert book["Change_Log"]["F2"].value == target
    assert book["Change_Log"]["F2"].data_type == "s"


def test_stale_preview_rejected():
    data = workbook()
    preview = preview_recode(data, "Data", "Sex", "f", "F")
    changed = apply_edits(data, "Data", [{"cell": "B2", "value": "X"}])
    with pytest.raises(ValueError, match="미리보기 이후"):
        apply_edits(changed, "Data", preview["edits"])


@pytest.mark.parametrize("source,target", [("", "F"), ("f", " "), ("f", "f")])
def test_invalid_mapping_rejected(source, target):
    with pytest.raises(ValueError):
        preview_recode(workbook(), "Data", "Sex", source, target)


def test_no_matches_and_unknown_column():
    assert preview_recode(workbook(), "Data", "Sex", "female", "F")["edits"] == []
    with pytest.raises(ValueError, match="열 이름"):
        preview_recode(workbook(), "Data", "Missing", "f", "F")
