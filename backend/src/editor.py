import re
from datetime import date, datetime
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

LOG_SHEET = "Change_Log"
LOG_HEADER = ["Changed_At", "Sheet", "Cell", "Column", "Before", "After", "Issue_Type"]

# 앞자리 0이 있는 값(예: 001)은 ID일 수 있으므로 숫자로 바꾸지 않는다
PLAIN_NUMBER = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")


# 입력한 글자를 엑셀에 넣을 값으로 변환: 숫자는 숫자로, 원래 날짜였던 셀은 날짜로
def convert_value(text, before):
    if text == "":
        return None

    if PLAIN_NUMBER.match(text):
        return float(text) if "." in text else int(text)

    if isinstance(before, (datetime, date)):
        parsed = pd.to_datetime(text.strip(), errors="coerce")
        if not pd.isna(parsed):
            return parsed.to_pydatetime()

    return text


def display(value):
    return "" if value is None else str(value)


# 원본 엑셀의 셀 값을 고치고, 고친 내역을 Change_Log 시트에 남긴다 (서식과 다른 시트는 그대로 유지)
def apply_edits(file_bytes, sheet, edits, changed_at=None):
    book = load_workbook(BytesIO(file_bytes))

    if sheet not in book.sheetnames:
        raise ValueError(f"시트 '{sheet}'가 파일에 없습니다.")

    ws = book[sheet]
    changed_at = changed_at or datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    if LOG_SHEET in book.sheetnames:
        log = book[LOG_SHEET]
    else:
        log = book.create_sheet(LOG_SHEET)
        log.append(LOG_HEADER)

    for edit in edits:
        cell = ws[edit["cell"]]
        if cell.row == 1:
            raise ValueError(f"{edit['cell']}: 머리글(1행)은 수정할 수 없습니다.")

        before = cell.value
        cell.value = convert_value(edit["value"], before)
        log.append([
            changed_at,
            sheet,
            cell.coordinate,
            display(ws.cell(row=1, column=cell.column).value),
            display(before),
            display(cell.value),
            edit.get("issue", ""),
        ])

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
